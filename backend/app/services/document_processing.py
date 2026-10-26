import io
import logging
import re
from dataclasses import dataclass

from docx import Document
from pypdf import PdfReader

logger = logging.getLogger(__name__)
MAX_PAGES = 500
MAX_EXTRACTED_CHARACTERS = 500_000
CHUNK_SIZE = 1_200
CHUNK_OVERLAP = 180
SUPPORTED_TEXT_EXTENSIONS = {".pdf", ".txt", ".docx"}


@dataclass(frozen=True)
class ProcessedDocument:
    status: str
    chunks: tuple[tuple[int | None, str], ...]


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _split(text: str, page_number: int | None) -> list[tuple[int | None, str]]:
    text = _normalize(text)
    output: list[tuple[int | None, str]] = []
    start = 0
    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        if end < len(text):
            boundary = max(text.rfind(". ", start + CHUNK_SIZE // 2, end), text.rfind("؟ ", start + CHUNK_SIZE // 2, end))
            if boundary > start:
                end = boundary + 1
        piece = text[start:end].strip()
        if piece:
            output.append((page_number, piece))
        if end == len(text):
            break
        start = max(start + 1, end - CHUNK_OVERLAP)
    return output


def process_document(file_name: str, content: bytes) -> ProcessedDocument:
    extension = "." + file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if extension not in SUPPORTED_TEXT_EXTENSIONS:
        return ProcessedDocument("unsupported", ())
    try:
        if extension == ".txt":
            pages = [(None, content.decode("utf-8"))]
        elif extension == ".docx":
            document = Document(io.BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            pages = [(None, "\n".join(parts))]
        else:
            reader = PdfReader(io.BytesIO(content), strict=False)
            if reader.is_encrypted:
                raise ValueError("Encrypted PDF is not supported")
            if len(reader.pages) > MAX_PAGES:
                raise ValueError("PDF page limit exceeded")
            pages = [(index + 1, page.extract_text() or "") for index, page in enumerate(reader.pages)]
    except Exception:
        logger.warning("Text extraction failed for an uploaded document.")
        return ProcessedDocument("failed", ())

    chunks: list[tuple[int | None, str]] = []
    character_count = 0
    truncated = False
    for page_number, page_text in pages:
        remaining = MAX_EXTRACTED_CHARACTERS - character_count
        if remaining <= 0:
            truncated = True
            break
        if len(page_text) > remaining:
            page_text = page_text[:remaining]
            truncated = True
        character_count += len(page_text)
        chunks.extend(_split(page_text, page_number))
    if not chunks:
        return ProcessedDocument("no_text", ())
    return ProcessedDocument("truncated" if truncated else "ready", tuple(chunks))
